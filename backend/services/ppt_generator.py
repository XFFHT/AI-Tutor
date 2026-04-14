from __future__ import annotations

import os
import re
import unicodedata
import hashlib
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Literal

from PIL import Image, ImageDraw, ImageFont
from matplotlib import font_manager
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

try:
    import win32com.client

    HAS_WIN32COM = True
except ImportError:
    HAS_WIN32COM = False

from services.latex_renderer import latex_renderer
from utils.config import TEMP_DIR


OPTION_PATTERN = re.compile(r"^\s*([A-H])[\.．、\)\]）]\s*(.*)$")
INLINE_MATH_PATTERN = re.compile(r"(?<!\\)\$(.*?)(?<!\\)\$", re.DOTALL)
EMU_PER_INCH = 914400
PIL_DPI = 96

COLOR_TITLE = (31, 91, 166)
COLOR_TEXT = (34, 34, 34)
COLOR_MUTED = (106, 119, 138)
COLOR_ANSWER = (20, 135, 84)
BLOCK_CACHE_VERSION = "v4-formula-size"


@dataclass
class InlineItem:
    kind: Literal["text", "formula"]
    content: str
    width: float
    height: float
    image_path: str | None = None


@dataclass
class LineLayout:
    items: list[InlineItem] = field(default_factory=list)
    width: float = 0.0
    height: float = 0.0


@dataclass
class FontProfile:
    stem: int
    option: int
    answer: int
    analysis: int


@dataclass
class BlockRender:
    path: str
    width: float
    height: float


class PPTGenerator:
    def __init__(self) -> None:
        self.slide_width_inches = 13.333
        self.slide_height_inches = 7.5
        self.page_left = 0.7
        self.page_right = 0.7
        self.content_width = self.slide_width_inches - self.page_left - self.page_right
        self.block_cache_dir = TEMP_DIR / "ppt_block_cache"
        self.block_cache_dir.mkdir(parents=True, exist_ok=True)

    def _inches(self, emu_value: int) -> float:
        return emu_value / EMU_PER_INCH

    def _normalize_text(self, text: str) -> str:
        text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
        text = text.replace("\u3000", " ")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    def contains_latex(self, text: str) -> bool:
        return bool(INLINE_MATH_PATTERN.search(text or ""))

    def split_inline_segments(self, text: str) -> list[tuple[Literal["text", "latex"], str]]:
        text = text or ""
        segments: list[tuple[Literal["text", "latex"], str]] = []
        last_end = 0

        for match in INLINE_MATH_PATTERN.finditer(text):
            start = match.start()
            if start > last_end:
                plain = text[last_end:start]
                if plain:
                    segments.append(("text", plain))

            formula = match.group(1)
            if formula:
                segments.append(("latex", formula.strip()))
            last_end = match.end()

        if last_end < len(text):
            tail = text[last_end:]
            if tail:
                segments.append(("text", tail))

        return segments

    def _is_cjk(self, char: str) -> bool:
        if not char:
            return False
        code = ord(char)
        return (
            0x4E00 <= code <= 0x9FFF
            or 0x3400 <= code <= 0x4DBF
            or 0x3040 <= code <= 0x30FF
            or 0xAC00 <= code <= 0xD7AF
            or unicodedata.east_asian_width(char) in {"W", "F"}
        )

    def _tokenize_text(self, text: str) -> list[str]:
        tokens: list[str] = []
        buffer: list[str] = []

        def flush() -> None:
            if buffer:
                tokens.append("".join(buffer))
                buffer.clear()

        for char in text:
            if char.isspace():
                flush()
                tokens.append(char)
                continue

            if self._is_cjk(char) or char in "，。；：！？、（）()【】[]<>《》“”‘’+-=*/,.:":
                flush()
                tokens.append(char)
                continue

            buffer.append(char)

        flush()
        return tokens

    def _resolve_font_path(self, bold: bool) -> str | None:
        candidates = [
            Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
            Path(r"C:\Windows\Fonts\simhei.ttf" if bold else r"C:\Windows\Fonts\simsun.ttc"),
            Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        ]
        for candidate in candidates:
            if candidate.exists():
                return str(candidate)

        preferred = "Microsoft YaHei"
        try:
            return font_manager.findfont(preferred, fallback_to_default=True)
        except Exception:
            return None

    @lru_cache(maxsize=128)
    def _get_pil_font(self, font_size: int, bold: bool) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
        px_size = max(int(round(font_size * PIL_DPI / 72)), 1)
        font_path = self._resolve_font_path(bold)
        if font_path:
            try:
                return ImageFont.truetype(font_path, px_size)
            except Exception:
                pass
        return ImageFont.load_default()

    def _text_metrics(self, text: str, font_size: int, bold: bool = False) -> tuple[float, float]:
        font = self._get_pil_font(font_size, bold)
        bbox = font.getbbox(text or " ")
        width_px = max(bbox[2] - bbox[0], 1)
        height_px = max(bbox[3] - bbox[1], 1)
        return width_px / PIL_DPI, height_px / PIL_DPI

    def _text_line_height(self, font_size: int, bold: bool = False) -> float:
        _, height = self._text_metrics("国Ay", font_size, bold)
        return max(height * 1.08, font_size * 1.08 / 72)

    def _to_px(self, inches: float) -> int:
        return max(int(round(inches * PIL_DPI)), 1)

    def _color_to_rgba(self, color: tuple[int, int, int], alpha: int = 255) -> tuple[int, int, int, int]:
        return (color[0], color[1], color[2], alpha)

    def _resample_filter(self):
        if hasattr(Image, "Resampling"):
            return Image.Resampling.LANCZOS
        return Image.LANCZOS

    @lru_cache(maxsize=2048)
    def _formula_metrics(self, formula: str, font_size: int) -> tuple[str, float, float]:
        render_size = max(int(round(font_size * 0.85)), 11)
        image_path = latex_renderer.render_to_image(formula, font_size=render_size)
        with Image.open(image_path) as image:
            width_px, height_px = image.size

        text_line_height = self._text_line_height(font_size, False)
        target_height = max(min(text_line_height * 0.82, font_size * 0.88 / 72), 0.14)
        aspect_ratio = width_px / max(height_px, 1)
        target_width = max(target_height * aspect_ratio, 0.14)
        return image_path, target_width, target_height

    def _new_line(self, font_size: int, bold: bool = False) -> LineLayout:
        return LineLayout(items=[], width=0.0, height=self._text_line_height(font_size, bold))

    def _append_text_token(
        self,
        lines: list[LineLayout],
        current_line: LineLayout,
        token: str,
        max_width: float,
        font_size: int,
        bold: bool,
    ) -> LineLayout:
        if token.isspace() and not current_line.items:
            return current_line

        token_width, token_height = self._text_metrics(token, font_size, bold)

        if token.strip() and current_line.items and current_line.width + token_width > max_width:
            lines.append(current_line)
            current_line = self._new_line(font_size, bold)

        if token.isspace() and not current_line.items:
            return current_line

        if current_line.items and current_line.items[-1].kind == "text":
            current_line.items[-1].content += token
            current_line.items[-1].width += token_width
            current_line.items[-1].height = max(current_line.items[-1].height, token_height)
        else:
            current_line.items.append(
                InlineItem(kind="text", content=token, width=token_width, height=token_height)
            )

        current_line.width += token_width
        current_line.height = max(current_line.height, token_height, self._text_line_height(font_size, bold))
        return current_line

    def _append_formula(
        self,
        lines: list[LineLayout],
        current_line: LineLayout,
        formula: str,
        max_width: float,
        font_size: int,
    ) -> LineLayout:
        image_path, formula_width, formula_height = self._formula_metrics(formula, font_size)

        if current_line.items and current_line.width + formula_width > max_width:
            lines.append(current_line)
            current_line = self._new_line(font_size)

        if formula_width > max_width:
            scale = max_width / formula_width
            formula_width = max_width
            formula_height *= scale

        current_line.items.append(
            InlineItem(
                kind="formula",
                content=formula,
                width=formula_width,
                height=formula_height,
                image_path=image_path,
            )
        )
        current_line.width += formula_width
        current_line.height = max(current_line.height, formula_height)
        return current_line

    def build_rich_lines(self, text: str, max_width: float, font_size: int, bold: bool = False) -> list[LineLayout]:
        normalized = self._normalize_text(text)
        raw_lines = normalized.split("\n") if normalized else [""]
        layouts: list[LineLayout] = []

        for raw_line in raw_lines:
            current_line = self._new_line(font_size, bold)

            if not raw_line:
                current_line.height = self._text_line_height(font_size, bold) * 0.8
                layouts.append(current_line)
                continue

            for segment_type, segment_content in self.split_inline_segments(raw_line):
                if segment_type == "text":
                    for token in self._tokenize_text(segment_content):
                        current_line = self._append_text_token(
                            layouts, current_line, token, max_width, font_size, bold
                        )
                else:
                    current_line = self._append_formula(
                        layouts, current_line, segment_content, max_width, font_size
                    )

            layouts.append(current_line)

        return layouts

    def estimate_rich_block_height(
        self,
        text: str,
        max_width: float,
        font_size: int,
        bold: bool = False,
        line_gap: float = 0.03,
    ) -> float:
        lines = self.build_rich_lines(text, max_width, font_size, bold)
        if not lines:
            return 0.0
        return sum(line.height for line in lines) + max(len(lines) - 1, 0) * line_gap

    def _add_text_shape(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        width: float,
        height: float,
        font_size: int,
        color: tuple[int, int, int],
        bold: bool = False,
        align: PP_ALIGN = PP_ALIGN.LEFT,
    ) -> None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(max(width, 0.2)), Inches(max(height, 0.2)))
        shape.fill.background()
        shape.line.fill.background()

        text_frame = shape.text_frame
        text_frame.clear()
        text_frame.word_wrap = False
        text_frame.margin_left = 0
        text_frame.margin_right = 0
        text_frame.margin_top = 0
        text_frame.margin_bottom = 0

        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.name = "Microsoft YaHei"
        run.font.color.rgb = RGBColor(*color)

    @lru_cache(maxsize=2048)
    def _render_rich_block_image(
        self,
        text: str,
        width: float,
        font_size: int,
        color: tuple[int, int, int],
        bold: bool,
        line_gap: float,
    ) -> BlockRender:
        lines = self.build_rich_lines(text, width, font_size, bold)
        if not lines:
            empty_path = self.block_cache_dir / "empty_block.png"
            if not empty_path.exists():
                Image.new("RGBA", (2, 2), (255, 255, 255, 0)).save(empty_path)
            return BlockRender(path=str(empty_path), width=0.02, height=0.02)

        horizontal_padding = 0.015
        vertical_padding = 0.015
        content_height = sum(line.height for line in lines) + max(len(lines) - 1, 0) * line_gap
        content_width = max((max((line.width for line in lines), default=0.0)), 0.2)
        image_width_inches = max(content_width + horizontal_padding * 2, 0.2)
        image_height_inches = max(content_height + vertical_padding * 2, 0.2)

        image = Image.new(
            "RGBA",
            (self._to_px(image_width_inches), self._to_px(image_height_inches)),
            (255, 255, 255, 0),
        )
        draw = ImageDraw.Draw(image)
        font = self._get_pil_font(font_size, bold)
        resample_filter = self._resample_filter()
        cursor_y = vertical_padding

        for line in lines:
            cursor_x = horizontal_padding
            if not line.items:
                cursor_y += line.height + line_gap
                continue

            for item in line.items:
                if item.kind == "text":
                    token_y = cursor_y + max((line.height - item.height) / 2, 0)
                    draw.text(
                        (self._to_px(cursor_x), self._to_px(token_y)),
                        item.content,
                        font=font,
                        fill=self._color_to_rgba(color),
                    )
                else:
                    token_y = cursor_y + max((line.height - item.height) / 2, 0)
                    with Image.open(item.image_path).convert("RGBA") as formula_image:
                        formula_image = formula_image.resize(
                            (self._to_px(item.width), self._to_px(item.height)),
                            resample=resample_filter,
                        )
                        image.alpha_composite(formula_image, (self._to_px(cursor_x), self._to_px(token_y)))

                cursor_x += item.width

            cursor_y += line.height + line_gap

        cache_key = f"{BLOCK_CACHE_VERSION}|{text}|{width:.3f}|{font_size}|{color}|{bold}|{line_gap:.3f}"
        cache_name = f"block_{hashlib.md5(cache_key.encode('utf-8')).hexdigest()}.png"
        output_path = self.block_cache_dir / cache_name
        if not output_path.exists():
            image.save(output_path)

        return BlockRender(
            path=str(output_path),
            width=content_width + horizontal_padding * 2,
            height=content_height + vertical_padding * 2,
        )

    def draw_rich_block(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        width: float,
        font_size: int,
        color: tuple[int, int, int] = COLOR_TEXT,
        bold: bool = False,
        line_gap: float = 0.03,
    ) -> float:
        block = self._render_rich_block_image(text, width, font_size, color, bold, line_gap)
        slide.shapes.add_picture(
            block.path,
            Inches(x),
            Inches(y),
            width=Inches(block.width),
            height=Inches(block.height),
        )
        return y + block.height

    def estimate_labeled_block_height(
        self,
        label: str,
        text: str,
        width: float,
        label_font_size: int,
        text_font_size: int,
        line_gap: float = 0.03,
    ) -> float:
        label_width, _ = self._text_metrics(label, label_font_size, True)
        value_width = max(width - label_width - 0.06, 1.0)
        value_height = self.estimate_rich_block_height(text, value_width, text_font_size, False, line_gap)
        return max(self._text_line_height(label_font_size, True), value_height)

    def draw_labeled_block(
        self,
        slide,
        label: str,
        text: str,
        x: float,
        y: float,
        width: float,
        label_font_size: int,
        text_font_size: int,
        label_color: tuple[int, int, int],
        text_color: tuple[int, int, int] = COLOR_TEXT,
        line_gap: float = 0.03,
    ) -> float:
        label_width, _ = self._text_metrics(label, label_font_size, True)
        label_height = self._text_line_height(label_font_size, True)
        self._add_text_shape(
            slide=slide,
            text=label,
            x=x,
            y=y,
            width=label_width + 0.04,
            height=label_height,
            font_size=label_font_size,
            color=label_color,
            bold=True,
        )

        value_x = x + label_width + 0.06
        value_width = max(width - label_width - 0.06, 1.0)
        value_bottom = self.draw_rich_block(
            slide=slide,
            text=text,
            x=value_x,
            y=y,
            width=value_width,
            font_size=text_font_size,
            color=text_color,
            bold=False,
            line_gap=line_gap,
        )
        return max(y + label_height, value_bottom)

    def split_question_content(self, content: str) -> tuple[str, list[tuple[str, str]]]:
        normalized = self._normalize_text(content)
        if not normalized:
            return "", []

        stem_lines: list[str] = []
        options: list[tuple[str, str]] = []
        current_option_letter: str | None = None
        current_option_lines: list[str] = []

        def flush_option() -> None:
            nonlocal current_option_letter, current_option_lines
            if current_option_letter is not None:
                option_text = "\n".join(line for line in current_option_lines if line is not None).strip()
                options.append((current_option_letter, option_text))
            current_option_letter = None
            current_option_lines = []

        for line in normalized.split("\n"):
            match = OPTION_PATTERN.match(line)
            if match:
                flush_option()
                current_option_letter = match.group(1)
                current_option_lines = [match.group(2).strip()]
                continue

            if current_option_letter is not None:
                current_option_lines.append(line.strip())
            else:
                stem_lines.append(line)

        flush_option()
        return "\n".join(stem_lines).strip(), options

    def _format_stem_for_display(self, text: str) -> str:
        text = self._normalize_text(text)
        if not text:
            return text

        replacements = [
            ("，且", "，\n且"),
            ("，则", "，\n则"),
            ("。则", "。\n则"),
            ("；", "；\n"),
        ]
        for source, target in replacements:
            text = text.replace(source, target)

        text = re.sub(r"(?<!\n)(\(\d+\)|（\d+）)", r"\n\1", text)
        text = re.sub(r"\n{2,}", "\n", text)
        return text.strip()

    def _format_analysis_for_display(self, text: str) -> str:
        text = self._normalize_text(text)
        if not text:
            return text

        text = re.sub(r"(?<!\n)(\(\d+\)|（\d+）)", r"\n\1", text)
        text = text.replace("。", "。\n")
        text = text.replace("；", "；\n")
        text = re.sub(r"\n{2,}", "\n", text)
        return text.strip()

    def choose_font_profile(
        self,
        stem: str,
        options: list[tuple[str, str]],
        answer: str,
        analysis: str,
        width: float,
        available_height: float,
    ) -> FontProfile:
        candidates = [
            FontProfile(stem=18, option=15, answer=15, analysis=13),
            FontProfile(stem=17, option=14, answer=14, analysis=12),
            FontProfile(stem=16, option=13, answer=13, analysis=11),
            FontProfile(stem=15, option=12, answer=12, analysis=10),
            FontProfile(stem=14, option=11, answer=11, analysis=10),
        ]

        for profile in candidates:
            total_height = 0.0

            if stem:
                total_height += self.estimate_rich_block_height(stem, width, profile.stem, False, 0.025)
                total_height += 0.06

            if options:
                for letter, option_text in options:
                    total_height += self.estimate_labeled_block_height(
                        label=f"{letter}.",
                        text=option_text,
                        width=width - 0.1,
                        label_font_size=profile.option,
                        text_font_size=profile.option,
                        line_gap=0.025,
                    )
                    total_height += 0.04

            total_height += self.estimate_labeled_block_height(
                label="【答案】",
                text=answer,
                width=width,
                label_font_size=profile.answer,
                text_font_size=profile.answer,
                line_gap=0.025,
            )
            total_height += 0.06

            total_height += self.estimate_labeled_block_height(
                label="【解析】",
                text=analysis,
                width=width,
                label_font_size=profile.analysis,
                text_font_size=profile.analysis,
                line_gap=0.025,
            )

            if total_height <= available_height:
                return profile

        return candidates[-1]

    def create_title_slide(self, prs: Presentation, main_title: str, subtitle_text: str) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        self._add_text_shape(
            slide=slide,
            text=main_title or "试卷讲解",
            x=0.8,
            y=2.15,
            width=self.slide_width_inches - 1.6,
            height=0.7,
            font_size=24,
            color=COLOR_TITLE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        self._add_text_shape(
            slide=slide,
            text=subtitle_text,
            x=0.8,
            y=3.0,
            width=self.slide_width_inches - 1.6,
            height=0.35,
            font_size=12,
            color=COLOR_MUTED,
            bold=False,
            align=PP_ALIGN.CENTER,
        )

    def create_question_slide(
        self,
        prs: Presentation,
        question_num: int,
        content: str,
        source: str,
        answer: str,
        analysis: str,
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        stem, options = self.split_question_content(content)
        stem = self._format_stem_for_display(stem)
        options = [(letter, self._format_stem_for_display(option_text)) for letter, option_text in options]
        answer = self._normalize_text(answer)
        analysis = self._format_analysis_for_display(analysis)
        source = self._normalize_text(source)
        body_width = min(self.content_width, 6.8)
        option_width = body_width - 0.15

        header_top = 0.16
        current_y = 0.26

        self._add_text_shape(
            slide=slide,
            text=f"第 {question_num} 题",
            x=0.8,
            y=header_top,
            width=self.slide_width_inches - 1.6,
            height=0.38,
            font_size=20,
            color=COLOR_TITLE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )

        current_y = 0.62
        if source:
            self._add_text_shape(
                slide=slide,
                text=f"来源：{source}",
                x=0.72,
                y=current_y,
                width=body_width,
                height=0.2,
                font_size=9,
                color=COLOR_MUTED,
                bold=False,
                align=PP_ALIGN.LEFT,
            )
            current_y += 0.24

        available_height = self.slide_height_inches - current_y - 0.22
        profile = self.choose_font_profile(
            stem=stem,
            options=options,
            answer=answer,
            analysis=analysis,
            width=body_width,
            available_height=available_height,
        )

        if stem:
            current_y = self.draw_rich_block(
                slide=slide,
                text=stem,
                x=self.page_left,
                y=current_y,
                width=body_width,
                font_size=profile.stem,
                color=COLOR_TEXT,
                bold=False,
                line_gap=0.025,
            )
            current_y += 0.05

        for letter, option_text in options:
            current_y = self.draw_labeled_block(
                slide=slide,
                label=f"{letter}.",
                text=option_text,
                x=self.page_left + 0.1,
                y=current_y,
                width=option_width,
                label_font_size=profile.option,
                text_font_size=profile.option,
                label_color=COLOR_TEXT,
                text_color=COLOR_TEXT,
                line_gap=0.025,
            )
            current_y += 0.03

        current_y += 0.01
        current_y = self.draw_labeled_block(
            slide=slide,
            label="【答案】",
            text=answer or "未提供",
            x=self.page_left,
            y=current_y,
            width=body_width,
            label_font_size=profile.answer,
            text_font_size=profile.answer,
            label_color=COLOR_ANSWER,
            text_color=COLOR_ANSWER,
            line_gap=0.025,
        )

        current_y += 0.04
        self.draw_labeled_block(
            slide=slide,
            label="【解析】",
            text=analysis or "未提供解析。",
            x=self.page_left,
            y=current_y,
            width=body_width,
            label_font_size=profile.analysis,
            text_font_size=profile.analysis,
            label_color=COLOR_MUTED,
            text_color=COLOR_TEXT,
            line_gap=0.025,
        )

    def add_animations_via_com(self, pptx_path: Path) -> bool:
        if not HAS_WIN32COM:
            return False

        app = None
        prs = None
        try:
            app = win32com.client.Dispatch("PowerPoint.Application")
            prs = app.Presentations.Open(os.path.abspath(str(pptx_path)), WithWindow=False)

            for slide in prs.Slides:
                for shape in slide.Shapes:
                    if shape.Name == "AnimatedAnswerShape":
                        slide.TimeLine.MainSequence.AddEffect(shape, 10, 0, 1)

            prs.Save()
            return True
        except Exception:
            return False
        finally:
            try:
                if prs is not None:
                    prs.Close()
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass

    def generate(
        self,
        questions: list,
        output_path: Path,
        title: str | None = None,
        use_animation: bool = True,
    ) -> Path:
        prs = Presentation()
        prs.slide_width = Inches(self.slide_width_inches)
        prs.slide_height = Inches(self.slide_height_inches)

        if title:
            main_title = title
        elif questions and questions[0].get("source"):
            main_title = str(questions[0]["source"])
        else:
            main_title = "试卷讲解"

        subtitle_text = f"共 {len(questions)} 道题目"
        self.create_title_slide(prs, main_title, subtitle_text)

        for idx, question in enumerate(questions, 1):
            self.create_question_slide(
                prs=prs,
                question_num=idx,
                content=str(question.get("content", "")),
                source=str(question.get("source", "")),
                answer=str(question.get("answer", "")),
                analysis=str(question.get("analysis", "")),
            )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))

        if use_animation:
            self.add_animations_via_com(output_path)

        return output_path


ppt_generator = PPTGenerator()
